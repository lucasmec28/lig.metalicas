"""Memória editável: equações OMML, tabela única de verificações e desenho comum."""
from io import BytesIO
from copy import deepcopy
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from .models import KGF,VERSION,STEELS,BOLTS
from .drawing import image_bytes

LINK='https://www.linkedin.com/in/lucas-oliveira-722723149/?isSelfProfile=true'
BRAND='Desenvolvido por LRO Soluções de engenharia LTDA.'


def number(x,places=2):
    return f'{x:,.{places}f}'.replace(',','§').replace('.',',').replace('§','.')


def display(v,unit):
    factor,label={'N':(KGF,'kgf'),'Nmm':(KGF*1000,'kgf·m'),'N/mm':(KGF,'kgf/mm')}.get(unit,(1,unit))
    return number(v/factor,3 if unit=='—' else 2),label


def hyperlink(p,label,url):
    rel=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel)
    r=OxmlElement('w:r');rp=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'225D86');rp.append(color);r.append(rp)
    t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)


def table(doc,headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for col,width in zip(t.columns,widths):col.width=Inches(width)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','bottom','left','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    props.append(borders)
    for i,label in enumerate(headers):t.rows[0].cells[i].text=label
    rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
    for row in rows:
        for cell,value in zip(t.add_row().cells,row):cell.text=str(value)
    for ri,row in enumerate(t.rows):
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
        for j,cell in enumerate(row.cells):
            cell.width=Inches(widths[j]);cell.vertical_alignment=1
            pr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE8F0' if ri==0 else ('F5F7F9' if ri%2==0 else 'FFFFFF'));pr.append(sh)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'65');el.set(qn('w:type'),'dxa');margins.append(el)
            pr.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
                if j>0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.size=Pt(9);r.bold=(ri==0)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t


# OMML nativo: frações, índices, expoentes e radicais estruturados.
def mr(text):
    r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=str(text);r.append(t);return r
def group(*items):
    out=[]
    for item in items:
        if isinstance(item,list):out.extend(item)
        else:out.append(mr(item) if isinstance(item,str) else item)
    return out
def slot(tag,items):
    e=OxmlElement('m:'+tag)
    for item in group(items):e.append(deepcopy(item))
    return e
def sub(base,idx):
    e=OxmlElement('m:sSub');e.append(slot('e',group(base)));e.append(slot('sub',group(idx)));return e
def power(base,exp='2'):
    e=OxmlElement('m:sSup');e.append(slot('e',group(base)));e.append(slot('sup',group(exp)));return e
def frac(a,b):
    e=OxmlElement('m:f');e.append(slot('num',group(a)));e.append(slot('den',group(b)));return e
def radical(items):
    e=OxmlElement('m:rad');pr=OxmlElement('m:radPr');de=OxmlElement('m:degHide');de.set(qn('m:val'),'1');pr.append(de);e.append(pr);e.append(slot('deg',[]));e.append(slot('e',group(items)));return e


def native_equation(doc,check,c):
    fy=sub('f','y');fu=sub('f','u');g1=sub('γ','a1');g2=sub('γ','a2');Ag=sub('A','g');An=sub('A','n')
    cid=check.id
    if cid=='bolts':
        alpha='0,45' if c.threads or c.bolt=='ASTM A307' else '0,56'
        expr=group(sub('F','Rd'),' = ',frac(group(alpha,'π',power(sub('d','b')),sub('f','ub')),group('4',g2)))
    elif cid.startswith('bearing'):
        expr=group(sub('F','Rd'),' = ',frac(group('min(1,2',sub('l','c'),'t',fu,'; 2,4',sub('d','b'),'t',fu,')'),g2))
    elif cid in ['plate_vy','plate_vu','plate_ny','plate_nu']:
        rupture=cid.endswith('u');shear=cid.startswith('plate_v')
        expr=group(sub('V' if shear else 'N','Rd'),' = ',frac(group('0,60' if shear else '',fu if rupture else fy,An if rupture else Ag),g2 if rupture else g1))
    elif cid.startswith('interaction'):
        expr=group('η = ',power(group('[n/2 + m]')),' + ',power('v'),' ≤ 1   (n < 0,2)')
        expr2=group('η = ',power(group('[n + 8m/9]')),' + ',power('v'),' ≤ 1   (n ≥ 0,2)')
        _math(doc,expr);_math(doc,expr2);return
    elif cid=='weld':
        expr=group(sub('q','max'),' = ',radical(group(power(group('(',frac('N','2h'),' + ',frac('3M',power('h')),')')),' + ',power(group('(',frac('V','2h'),')')))))
        _math(doc,expr)
        expr=group(sub('q','Rd'),' = ',frac(group('0,60',sub('f','w'),'w'),group(radical('2'),sub('γ','w2'))))
    elif cid=='plate_mu':expr=group(sub('M','Rd'),' = ',frac(group(fu,sub('Z','n')),g2))
    elif cid=='plate_ltb':expr=group(sub('M','Rd'),' = ',frac(sub('M','n'),g1),' ; ',sub('M','n'),' = min(',sub('M','p'),'; ',sub('C','b'),'[1,52 − 0,274λ',frac(fy,'E'),']',fy,'W)')
    elif cid=='beam_vm':expr=group('η = ',frac('M',group(sub('M','BC,Rd'),' + ',sub('V','AB,Rd'),'ℓ')),' + ',frac('N',sub('N','web,Rd')),' ≤ 1')
    elif cid=='support_punch':expr=group(sub('t','p'),' ≤ ',frac(group(sub('t','s'),sub('f','u,s')),group(sub('f','y,p'),g2)))
    elif cid=='support_shear':expr=group(sub('R','Rd'),' = min(',frac(group('1,2',fy,'h',sub('t','s')),g1),'; ',frac(group('1,2',fu,'h',sub('t','s')),g2),')')
    elif cid=='support_web_y':expr=group(sub('R','Rd'),' = ',frac(group('1,10(2,5k + ',sub('l','n'),')',fy,sub('t','w')),g1))
    elif cid=='support_flange':expr=group(sub('R','Rd'),' = ',frac(group('0,5 × 6,25',power(sub('t','f')),fy),g1))
    elif cid.startswith('block'):
        expr=group(sub('R','Rd'),' = ',frac(group('min(0,60',fu,sub('A','nv'),' + ',fu,sub('A','nt'),'; 0,60',fy,sub('A','gv'),' + ',fu,sub('A','nt'),')'),g2))
        if cid=='block_plate':
            _math(doc,expr)
            expr=group('η = ',power(frac('V',sub('R','v'))),' + ',power(frac('N',sub('R','n'))),' ≤ 1')
    elif cid=='base_plate':expr=group(sub('q','Rd'),' = min(',frac(group('0,60',fy,sub('t','p')),g1),'; ',frac(group('0,60',fu,sub('t','p')),g2),')')
    elif cid=='beam_v':expr=group(sub('V','Rd'),' = ',frac(group('0,60',fy,sub('h','web'),sub('t','w')),g1))
    else:expr=group(sub('N','Rd'),' = min(',frac(group(fy,Ag),g1),'; ',frac(group(fu,An),g2),')')
    _math(doc,expr)


def _math(doc,expr):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(4)
    m=OxmlElement('m:oMath')
    for item in expr:m.append(deepcopy(item))
    p._p.append(m)


def create_report(c,r,detailed=False):
    doc=Document();s=doc.sections[0]
    s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=s.bottom_margin=Inches(.6);s.left_margin=s.right_margin=Inches(.7)
    for name in ['Normal','Title','Heading 1','Heading 2']:
        st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
    normal=doc.styles['Normal'];normal.font.size=Pt(9.5);normal.paragraph_format.space_after=Pt(5)
    doc.styles['Title'].font.size=Pt(20)
    doc.styles['Heading 1'].font.size=Pt(12);doc.styles['Heading 2'].font.size=Pt(10)
    doc.core_properties.author='LRO Soluções de engenharia LTDA.'
    doc.core_properties.title='Memória de cálculo de ligação com chapa simples'
    doc.add_paragraph('Ligação com chapa simples',style='Title')
    doc.add_paragraph(c.project)
    doc.add_paragraph(f'{c.beam.name} → {c.support.name} | '+('Alma da viga de apoio' if c.kind=='beam_web' else 'Mesa do pilar alinhada à alma'))
    p=doc.add_paragraph();p.add_run(r.status).bold=True
    if r.governing:
        p.add_run(f' · índice determinante {number(r.governing.ratio,3)} · {r.governing.name}')
    pending=[i for i in r.issues if i.severity in ('pending','error')]
    if pending:
        doc.add_paragraph('Conclusão condicionada às pendências abaixo. Este documento não aprova os estados limite ainda não verificados.')
        for i in pending:doc.add_paragraph(i.text)
    elif any(not x.passed for x in r.checks):
        doc.add_paragraph('A ligação não atende às verificações indicadas no quadro resumo. Revisar a geometria ou as especificações antes de utilizar o detalhe.')
    else:doc.add_paragraph('As verificações locais incluídas nesta versão atendem para a geometria e as hipóteses registradas. A análise global dos membros e os requisitos globais de integridade estrutural permanecem no projeto da estrutura.')
    doc.add_picture(BytesIO(image_bytes(c)),width=Inches(7.05))
    doc.add_paragraph('Dimensões e materiais adotados',style='Heading 1')
    geom=[('Chapa',f'{number(c.width)} × {number(c.hp)} × {number(c.tp,4)} mm',c.plate_steel),('Parafusos',f'{c.n} × Ø {number(c.db,3)} mm',c.bolt),('Furos',f'Ø {number(c.dh,4)} mm', 'Padrão; '+('broca' if c.drilled else 'desconto líquido +2 mm')),('Soldas',f'2 filetes de {number(c.weld)} mm; L = {number(c.hp)} mm',f'fw = {number(c.fw,0)} MPa'),('Posições',f'a = {c.a:g}; g = {c.gap:g}; p = {c.pitch:g}; eᵥ = eₕ = {c.edge_v:g}' if c.edge_v==c.edge_h else f'a = {c.a:g}; g = {c.gap:g}; p = {c.pitch:g}; eᵥ = {c.edge_v:g}; eₕ = {c.edge_h:g}', 'mm'),('Viga apoiada',f'd/bf/tw/tf = {c.beam.d:g}/{c.beam.bf:g}/{c.beam.tw:g}/{c.beam.tf:g}',c.beam_steel),('Apoio',f'd/bf/tw/tf = {c.support.d:g}/{c.support.bf:g}/{c.support.tw:g}/{c.support.tf:g}',c.support_steel)]
    table(doc,['Componente','Dimensões','Material ou especificação'],geom,[1.1,3.25,2.75])
    mats=[]
    for label,name in [('Chapa',c.plate_steel),('Viga',c.beam_steel),('Apoio',c.support_steel)]:
        st=STEELS[name];mats.append(f'{label}: fy = {st.fy:g} MPa; fu = {st.fu:g} MPa')
    doc.add_paragraph('; '.join(mats)+'. E = 200 000 MPa; γa1 = 1,10; γa2 = γw2 = 1,35.')
    doc.add_paragraph('Esforços e hipóteses',style='Heading 1')
    doc.add_paragraph(f'Entrada já majorada: V = {number(c.V/KGF)} kgf ({number(c.V/1000,3)} kN); N = {number(c.N/KGF)} kgf ({number(c.N/1000,3)} kN), tração positiva. Não se aplicam novos coeficientes de ações.')
    if r.minimum_factor>1:
        doc.add_paragraph(f'Verificação adicional do mínimo de 45 kN conforme NBR 8800, 6.1.5.2, na direção da resultante informada: V = {number(c.V*r.minimum_factor/KGF)} kgf; N = {number(c.N*r.minimum_factor/KGF)} kgf. Os resultados governantes abaixo incluem essa verificação.')
    doc.add_paragraph(f'Encontro a 90°; uma coluna de parafusos; soldagem em oficina; e = a = {number(c.a)} mm. Excentricidade vertical de N: eN = {number(c.eccentric_n)} mm. M = |V|a + |N·eN|. Contenção longitudinal da viga: '+('confirmada pelo usuário.' if c.restrained else 'não confirmada.'))
    cope_label={'none':'sem recorte','top':'mesa superior','both':'mesas superior e inferior'}[c.cope]
    doc.add_paragraph(f'Posição da chapa: topo a {number(c.plate_top)} mm do topo da viga. Desnível entre vigas: {number(c.beam_level)} mm. Recorte: {cope_label}; dimensões superior/inferior/comprimento = {c.coped_top:g}/{c.coped_bottom:g}/{c.cope_length if c.cope!="none" else 0:g} mm. Folga de montagem = {c.clearance:g} mm; raio do envelope de montagem = {c.tool_radius:g} mm. Borda da solda reforçada: '+('sim.' if c.reinforced_weld else 'não.'))
    if c.notes:doc.add_paragraph(c.notes)
    doc.add_paragraph('Referências e limites de aplicação',style='Heading 1')
    doc.add_paragraph('ABNT NBR 8800:2024, errata 2025: 5.2.4; 5.7; 6.1.5.2; 6.2; 6.3; 6.5 e Anexo A. AISC Companion v16.0, Volume 1, P901-23W: exemplos II.A-17B, II.A-18 e II.A-19B. Carini, Ligações Flexíveis: p.32–43. Procedimentos complementares AISC/SCI identificados por verificação, com resistências e coeficientes explicitados.')
    doc.add_paragraph('Versão 0.1: tração axial e cortante no plano; uma viga; sem fadiga, atrito, ação cíclica, incêndio ou dimensionamento de enrijecedores. Recortes e apoio em alma sob N permanecem condicionados às verificações externas indicadas. A força mínima é avaliada pela resultante; combinações de ações e exigências globais de integridade são externas.')
    doc.add_paragraph('Resumo das verificações',style='Heading 1')
    checkrows=[]
    for x in r.checks:
        sd,unit=display(x.demand,x.unit);rd,_=display(x.resistance,x.unit)
        checkrows.append((x.name,sd,rd,unit,number(x.ratio,3),'OK' if x.passed else 'NÃO'))
    table(doc,['Verificação','Sd','Rd','Unid.','Índice','Atende'],checkrows,[2.95,1,1,.65,.8,.7])
    doc.add_paragraph('Índices de interação são avaliados diretamente contra 1,0 e não representam um multiplicador linear da carga. Valores internos não são arredondados para decidir o atendimento.')
    doc.add_paragraph('Memória de cálculo',style='Heading 1')
    doc.add_paragraph('As substituições usam N, mm e MPa; os resultados principais são apresentados em kgf e kgf·m. Conversão: 1 kgf = 9,80665 N. Variáveis geométricas seguem o desenho. Na versão compacta, desenvolvem-se as verificações determinantes por componente.')
    if detailed:selected=r.checks
    else:
        groups=[['bolts'],['bearing_plate','bearing_beam'],['plate_ltb','plate_mu','interaction_y','interaction_u','block_plate','block_plate_u'],['beam_v','beam_n','block_beam_u','beam_vm'],['weld','base_plate'],['support_shear','support_punch','support_web_y','support_flange']]
        selected=[]
        for groupids in groups:
            candidates=[x for x in r.checks if x.id in groupids]
            if candidates:selected.append(max(candidates,key=lambda x:x.ratio))
    for x in selected:
        start=len(doc.paragraphs)
        doc.add_paragraph(x.name,style='Heading 2')
        native_equation(doc,x,c)
        doc.add_paragraph(x.substitution)
        sd,u=display(x.demand,x.unit);rd,_=display(x.resistance,x.unit)
        doc.add_paragraph(f'Resultado: Sd = {sd} {u}; Rd = {rd} {u}; índice = {number(x.ratio,3)}. '+('Atende.' if x.passed else 'Não atende.'))
        doc.add_paragraph(x.variables+' Referência: '+x.reference)
        for paragraph in doc.paragraphs[start:-1]:
            paragraph.paragraph_format.keep_with_next=True
    p=s.footer.paragraphs[0];p.paragraph_format.space_after=Pt(0)
    r0=p.add_run(BRAND+' ');r0.font.size=Pt(7)
    hyperlink(p,'LinkedIn de Lucas Oliveira',LINK)
    p=s.footer.add_paragraph(f'LRO Ligações v{VERSION} · ');p.paragraph_format.space_after=Pt(0)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
    for r0 in p.runs:r0.font.size=Pt(7)
    buf=BytesIO();doc.save(buf);return buf.getvalue()
